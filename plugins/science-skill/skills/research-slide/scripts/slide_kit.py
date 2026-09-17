#!/usr/bin/env python
"""slide_kit — a tiny toolkit for building editable Chinese research / defense
slides with python-pptx, in the house "research-highlight" style.

Why this exists
---------------
Every slide in this style repeats the same chrome: a navy statement banner,
□/■ section headers, a dashed column divider, gray method tags, red-italic
citations, boxed conclusions, and CJK-aware text runs with subscripts. Writing
that by hand per slide is error-prone (CJK fonts, subscript baselines, word-wrap
that silently breaks a header into two lines). This module captures it once so a
new slide is a few dozen readable lines.

Everything is EDITABLE in PowerPoint: text is real text (retypable), chrome is
native shapes, and only the figure panels are placed images.

Quick start
-----------
    from slide_kit import Deck, run, NAVY, BLACK, RED, LBLUE, ALIGN

    deck = Deck()                      # 16:9, 13.333 x 7.5 in
    s = deck.slide()

    s.corner_accent()
    s.header(0.42, 0.30, 12.5, [run("研究思路：零磁场实现分数量子反常霍尔效应",
                                     size=16, bold=True, color=BLACK)], mark="□")
    s.banner([run("拓扑平带 + 强关联相互作用  →  ", size=15, bold=True, color=WHITE),
              run("零磁场分数量子化", size=15, bold=True, color=LBLUE)], y=1.0)
    s.vdivider(6.66, 1.05, 7.2)
    w, h = s.picture("panels/def.png", 7.05, 1.05, h=1.8)
    s.cite(7.0, 6.9, 6.0, s.citation_runs("Park et al.", "Nature", "622",
                                          ", 74 (2023)", "共同一作"))
    deck.save("out.pptx")

All coordinates are INCHES. See references/design-system.md for the palette,
the two layout templates, and coordinate tables.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from PIL import Image

# --------------------------------------------------------------------- fonts
# CJK: PingFang SC ships on macOS. On other machines swap to a present CJK face
# (e.g. "Microsoft YaHei", "Noto Sans CJK SC"). SERIF is used for Latin text,
# math variables, and citations so they read like a journal figure.
CJK = "PingFang SC"
SERIF = "Times New Roman"

# --------------------------------------------------------------------- palette
NAVY     = RGBColor(0x1F, 0x38, 0x64)   # statement banner, conclusion accent, ■ marks
LBLUE    = RGBColor(0x9E, 0xC6, 0xF0)   # accent word inside the navy banner
BLUE     = RGBColor(0x2E, 0x75, 0xD6)   # corner accent + dashed divider
BLACK    = RGBColor(0x1A, 0x1A, 0x1A)   # header text
BODY     = RGBColor(0x33, 0x39, 0x40)   # body / highlight text
GRAY     = RGBColor(0x6B, 0x70, 0x78)   # captions, method tags
RED      = RGBColor(0xC0, 0x00, 0x00)   # citations, 实验挑战 emphasis
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
HAIRLINE = RGBColor(0xC7, 0xCE, 0xD6)   # thin rules / panel borders

# re-exported so callers don't import from pptx directly
ALIGN = PP_ALIGN
ANCHOR = MSO_ANCHOR


def run(text, size=14, bold=False, italic=False, color=BODY,
        sub=False, sup=False, serif=False):
    """One styled text run. `serif=True` uses Times (Latin/math/citations);
    otherwise PingFang (CJK). `sub`/`sup` set the subscript/superscript baseline
    — use these for ρ_yx, R_xx, e², etc. For chemical formulas prefer the
    unicode subscripts ₂₃₄ inside a normal run (MnBi₂Te₄), which stay on the
    baseline and kern better than manual sub runs."""
    fam = SERIF if serif else CJK
    return dict(t=text, size=size, bold=bold, italic=italic, color=color,
                sub=sub, sup=sup, latin=fam, ea=fam)


def _apply(r, spec):
    f = r.font
    f.size = Pt(spec['size']); f.bold = spec['bold']; f.italic = spec['italic']
    f.color.rgb = spec['color']
    rPr = r._r.get_or_add_rPr()
    latin_el = rPr.get_or_add_latin(); latin_el.set('typeface', spec['latin'])
    ea_el = rPr.find(qn('a:ea'))
    if ea_el is None:
        ea_el = rPr.makeelement(qn('a:ea'), {}); latin_el.addnext(ea_el)
    ea_el.set('typeface', spec['ea'])
    if spec.get('sub'):
        rPr.set('baseline', '-22000')
    elif spec.get('sup'):
        rPr.set('baseline', '30000')


def _paras(paras):
    """Accept either a single paragraph (list of run-dicts) or a list of
    paragraphs (list of lists). Returns a list of paragraphs."""
    if paras and isinstance(paras[0], dict):
        return [paras]
    return paras


class Slide:
    def __init__(self, slide, width, height):
        self.s = slide
        self.W = width
        self.H = height

    # ---------------------------------------------------------------- internals
    def _no_shadow(self, sp):
        try:
            sp.shadow.inherit = False
        except Exception:
            pass

    # ---------------------------------------------------------------- primitives
    def rrect(self, x, y, w, h, fill, line=None, line_w=1.0, radius=0.10):
        sp = self.s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                     Inches(x), Inches(y), Inches(w), Inches(h))
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
        if fill is None:
            sp.fill.background()
        else:
            sp.fill.solid(); sp.fill.fore_color.rgb = fill
        if line is None:
            sp.line.fill.background()
        else:
            sp.line.color.rgb = line; sp.line.width = Pt(line_w)
        self._no_shadow(sp)
        return sp

    def rect(self, x, y, w, h, fill=None, line=None, line_w=1.0):
        sp = self.s.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                     Inches(x), Inches(y), Inches(w), Inches(h))
        if fill is None:
            sp.fill.background()
        else:
            sp.fill.solid(); sp.fill.fore_color.rgb = fill
        if line is None:
            sp.line.fill.background()
        else:
            sp.line.color.rgb = line; sp.line.width = Pt(line_w)
        self._no_shadow(sp)
        return sp

    def line(self, x0, y0, x1, y1, color, w=1.5, dash=None):
        c = self.s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                        Inches(x0), Inches(y0), Inches(x1), Inches(y1))
        c.line.color.rgb = color; c.line.width = Pt(w)
        if dash:
            ln = c.line._get_or_add_ln()
            ln.append(ln.makeelement(qn('a:prstDash'), {'val': dash}))
        self._no_shadow(c)
        return c

    def vdivider(self, x, y0, y1, color=BLUE, w=2.0, dash="dash"):
        """The signature dashed vertical column separator."""
        return self.line(x, y0, x, y1, color, w=w, dash=dash)

    def hrule(self, x0, x1, y, color=HAIRLINE, w=1.0):
        return self.line(x0, y, x1, y, color, w=w)

    def text(self, x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             line_spacing=1.12, wrap=True, pad=0.0):
        """Place a text box. `paras` is a paragraph (list of run-dicts) or a
        list of paragraphs. `wrap=False` keeps a short label on ONE line even if
        the box is a hair narrow — set it on headers/labels so a stray wrap
        doesn't split '(a)' into two lines."""
        tb = self.s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = wrap
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = anchor
        for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
            setattr(tf, m, Inches(pad))
        for i, para in enumerate(_paras(paras)):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.line_spacing = line_spacing
            p.space_before = Pt(0); p.space_after = Pt(2)
            for spec in para:
                r = p.add_run(); r.text = spec['t']; _apply(r, spec)
        return tb

    def picture(self, path, x, y, h=None, w=None):
        """Place an image keeping aspect ratio; give exactly one of h/w
        (inches). Returns (w, h) so you can lay out neighbours. Panels should be
        pre-trimmed to tight bounds (see crop_panels.py)."""
        im = Image.open(path); ar = im.width / im.height
        if h is None and w is None:
            raise ValueError("give h or w")
        if h is not None:
            w = h * ar
        else:
            h = w / ar
        self.s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
        return w, h

    # ---------------------------------------------------------------- components
    def corner_accent(self, color=BLUE):
        """Small skewed blue tab in the very top-left corner — the deck's mark."""
        sp = self.s.shapes.add_shape(MSO_SHAPE.PARALLELOGRAM, Inches(-0.15),
                                     Inches(-0.10), Inches(1.02), Inches(0.40))
        try:
            sp.adjustments[0] = 0.6
        except Exception:
            pass
        sp.fill.solid(); sp.fill.fore_color.rgb = color; sp.line.fill.background()
        self._no_shadow(sp)
        return sp

    def banner(self, runs, x=None, y=0.26, w=None, h=0.60, fill=NAVY, radius=0.12):
        """Full-width navy statement banner with centered white text. Use it for
        the ONE-LINE thesis / core idea. Highlight a keyword by giving that run
        color=LBLUE."""
        x = 0.40 if x is None else x
        w = (self.W - 2 * x) if w is None else w
        self.rrect(x, y, w, h, fill, radius=radius)
        self.text(x, y, w, h, runs, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    def header(self, x, y, w, runs, mark="■", mark_color=NAVY, mark_size=13.5):
        """Section header: a □ (approach) or ■ (finding/evidence) marker plus
        bold text. Pass `mark=None` for no marker."""
        pre = [run(mark + "  ", size=mark_size, bold=True, color=mark_color)] if mark else []
        self.text(x, y, w, 0.6, [pre + runs], align=PP_ALIGN.LEFT,
                  anchor=MSO_ANCHOR.TOP, line_spacing=1.1)

    def topicbar(self, x, y, w, runs, fill=BLUE, h=0.50, radius=0.14):
        """Filled title bar (white text). Available, but prefer `header` with □/■
        for a cleaner look — large saturated bars tend to read as heavy."""
        self.rrect(x, y, w, h, fill, radius=radius)
        self.text(x + 0.12, y, w - 0.24, h, runs, align=PP_ALIGN.LEFT,
                  anchor=MSO_ANCHOR.MIDDLE)

    def tag(self, x, y, w, text_str, size=10.5, color=GRAY):
        """Small gray sub-label under a header (e.g. the 光/电 methods used)."""
        self.text(x, y, w, 0.24, [run(text_str, size=size, color=color)],
                  align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)

    def highlight(self, x, y, w, runs, label="研究亮点：", label_color=None,
                  size=12, h=1.4):
        """A labelled highlight paragraph: bold colored label + body runs."""
        label_color = NAVY if label_color is None else label_color
        pre = [run(label, size=size, bold=True, color=label_color)] if label else []
        self.text(x, y, w, h, [pre + runs], align=PP_ALIGN.LEFT, line_spacing=1.16)

    def caption(self, x, y, w, text_str, size=10, color=GRAY):
        self.text(x, y, w, 0.24, [run(text_str, size=size, color=color)],
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    def cite(self, x, y, w, runs, align=PP_ALIGN.LEFT):
        """Place a citation line (build the runs with `citation_runs`)."""
        self.text(x, y, w, 0.34, [runs], align=align, anchor=MSO_ANCHOR.MIDDLE)

    def citation_runs(self, first, journal, vol, rest, role=None, size=11,
                      color=RED):
        """Journal-style citation runs: FirstAuthor et al., *Journal* **Vol**,
        Page (Year)  role. Journal italic, volume bold, all serif. e.g.
        citation_runs("Park et al.", "Nature", "622", ", 74 (2023)", "共同一作")"""
        rs = [run(first + " ", size=size, color=color, serif=True),
              run(journal + " ", size=size, italic=True, color=color, serif=True),
              run(vol, size=size, bold=True, color=color, serif=True),
              run(rest, size=size, color=color, serif=True)]
        if role:
            rs.append(run("  " + role, size=size, color=color))
        return rs

    def box(self, x, y, w, h, paras, line=BLACK, line_w=1.4, fill=WHITE, pad=0.16):
        """Bordered conclusion box (right column of the single-story template).
        `paras` is a list of paragraphs."""
        self.rect(x, y, w, h, fill=fill, line=line, line_w=line_w)
        self.text(x + pad, y + pad, w - 2 * pad, h - 2 * pad, paras,
                  align=PP_ALIGN.LEFT, line_spacing=1.16)

    def conclusion(self, runs, y, x=0.40, w=None, rule=True, rule_gap=0.08):
        """Full-width synthesizing conclusion line (thesis→evidence template):
        an optional hairline, then centered bold text. Prefix a run with '▶ '."""
        w = (self.W - 2 * x) if w is None else w
        if rule:
            self.hrule(x, x + w, y - rule_gap)
        self.text(x, y, w, 0.52, [runs], align=PP_ALIGN.CENTER,
                  anchor=MSO_ANCHOR.MIDDLE)


class Deck:
    """A presentation. `slide()` returns a blank Slide; `save()` writes .pptx."""

    def __init__(self, width=13.333, height=7.5):
        self.prs = Presentation()
        self.prs.slide_width = Inches(width)
        self.prs.slide_height = Inches(height)
        self.W = width
        self.H = height

    def slide(self):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])  # blank
        return Slide(s, self.W, self.H)

    def save(self, path):
        self.prs.save(path)
        return path
